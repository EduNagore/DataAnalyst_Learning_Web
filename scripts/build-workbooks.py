"""Genera los libros Excel de práctica (public/workbooks/*.xlsx) a partir de Lumen.

Cada ejercicio produce dos libros: el de trabajo (`<id>.xlsx`) y el de solución
(`<id>-solucion.xlsx`). Ambos incluyen una hoja de autocomprobación que marca ✔/✘
comparando la respuesta con el valor esperado, calculado aquí con pandas de forma
independiente de las fórmulas del libro.

Las fórmulas del libro de solución usan funciones clásicas (SUMIFS, COUNTIF, INDEX/MATCH,
SUMPRODUCT) para que se calculen en cualquier versión de Excel y se puedan verificar fuera
de Excel (tests/workbooks); la columna «Fórmula moderna» muestra la equivalente con las
funciones dinámicas que se trabajan en la lección.

Uso: python scripts/build-workbooks.py
"""

from __future__ import annotations

import sys
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "public" / "data" / "lumen"
OUT = ROOT / "public" / "workbooks"
SEED = 20261008
N_ORDERS = 2000

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF3C4")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
KO_FILL = PatternFill("solid", fgColor="FFC7CE")
WRAP = Alignment(wrap_text=True, vertical="top")


@dataclass
class Task:
    question: str
    expected: float | str
    classic: str  # fórmula clásica (inglés, comas) para el libro de solución
    modern: str  # fórmula moderna sugerida (texto, en español)
    functions: str  # funciones sugeridas (ES / EN)
    tol: float = 0.005


@dataclass
class Exercise:
    id: str
    title: str
    intro: list[str]
    tasks: list[Task] = field(default_factory=list)
    kind: str = "tasks"  # "tasks" | "model"


# ---------------------------------------------------------------------------
# Datos
# ---------------------------------------------------------------------------


def load_sample() -> tuple[pd.DataFrame, pd.DataFrame]:
    orders = pd.read_parquet(DATA / "orders.parquet")
    items = pd.read_parquet(DATA / "order_items.parquet")
    products = pd.read_parquet(DATA / "products.parquet")[["product_id", "cost"]]
    customers = pd.read_parquet(DATA / "customers.parquet")

    sample_ids = orders.sample(N_ORDERS, random_state=SEED)["order_id"]
    it = items[items["order_id"].isin(sample_ids)].merge(products, on="product_id")
    it["importe"] = it["quantity"] * it["unit_price"]
    it["coste"] = it["quantity"] * it["cost"]
    agg = it.groupby("order_id")[["importe", "coste"]].sum().round(2).reset_index()

    p = orders[orders["order_id"].isin(sample_ids)].merge(agg, on="order_id", how="left")
    p = p.sort_values("order_id").reset_index(drop=True)
    p["importe"] = p["importe"].fillna(0.0)
    p["coste"] = p["coste"].fillna(0.0)
    p = p[
        [
            "order_id",
            "customer_id",
            "order_date",
            "channel",
            "discount_pct",
            "status",
            "importe",
            "coste",
        ]
    ]

    # Fichas de cliente: se omiten ~5 % a propósito (clientes «sin ficha»)
    c = customers[customers["customer_id"].isin(p["customer_id"])].copy()
    rng = np.random.default_rng(SEED)
    c = c[rng.random(len(c)) > 0.05]
    # El cliente del pedido de mayor importe debe tener ficha (tarea de BUSCARX)
    top_customer = p.loc[p["importe"].idxmax(), "customer_id"]
    if top_customer not in set(c["customer_id"]):
        c = pd.concat([c, customers[customers["customer_id"] == top_customer]])
    c = c[["customer_id", "segment", "region", "acquisition_channel"]]
    c = c.sort_values("customer_id").reset_index(drop=True)
    return p, c


# ---------------------------------------------------------------------------
# Escritura de hojas
# ---------------------------------------------------------------------------


def write_table(ws, df: pd.DataFrame, name: str, date_cols: tuple[str, ...] = ()) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else v for v in row])
    for cell in ws[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    for j, col in enumerate(df.columns, start=1):
        letter = ws.cell(row=1, column=j).column_letter
        ws.column_dimensions[letter].width = max(12, min(22, len(col) + 4))
        if col in date_cols:
            for r in range(2, len(df) + 2):
                ws.cell(row=r, column=j).number_format = "yyyy-mm-dd"
        elif col in ("importe", "coste"):
            for r in range(2, len(df) + 2):
                ws.cell(row=r, column=j).number_format = "#,##0.00"
        elif col == "discount_pct":
            for r in range(2, len(df) + 2):
                ws.cell(row=r, column=j).number_format = "0%"
    last_col = ws.cell(row=1, column=len(df.columns)).column_letter
    tab = Table(displayName=name, ref=f"A1:{last_col}{len(df) + 1}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A2"


def write_instructions(wb: Workbook, ex: Exercise, solution: bool) -> None:
    ws = wb.active
    ws.title = "Instrucciones"
    ws.column_dimensions["A"].width = 110
    lines = [f"Data Analyst Academy · {ex.title}" + (" · SOLUCIÓN" if solution else ""), ""]
    lines += ex.intro
    lines += [
        "",
        "Cómo se corrige: la hoja «Comprobacion» compara tu respuesta con el valor esperado y marca ✔ o ✘.",
        "Datos: muestra reproducible del dataset sintético Lumen (semilla fija). No son datos reales.",
        "Nombres de función: se muestran en español / inglés. El archivo guarda siempre los nombres en inglés.",
        "Separadores: en un Excel en español de España los argumentos se separan con «;» y el decimal es «,».",
    ]
    if solution:
        lines += [
            "",
            "En este libro las respuestas usan funciones clásicas para calcularse en cualquier versión;",
            "la columna «Fórmula moderna» de la hoja «Trabajo» muestra la equivalente con funciones dinámicas.",
        ]
    for i, text in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = WRAP
        if i == 1:
            c.font = Font(bold=True, size=14)


def write_check_sheet(wb: Workbook, labels: list[str], refs: list[str], expected: list, tols: list):
    ws = wb.create_sheet("Comprobacion")
    ws["A1"], ws["A1"].font = "Autocomprobación", Font(bold=True, size=14)
    headers = ["Nº", "Qué se comprueba", "Tu respuesta", "Esperado", "Resultado"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for i, (lab, ref, exp, tol) in enumerate(zip(labels, refs, expected, tols, strict=True)):
        r = 4 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=lab).alignment = WRAP
        ws.cell(row=r, column=3, value=f'=IF({ref}="","",{ref})')
        ws.cell(row=r, column=4, value=exp)
        ws.cell(row=r, column=5).value = (
            f'=IF(C{r}="","—",IF(ISNUMBER(D{r}),IF(ISNUMBER(C{r}),'
            f'IF(ABS(C{r}-D{r})<={tol},"✔","✘"),"✘"),IF(C{r}=D{r},"✔","✘")))'
        )
    end = 3 + len(labels)
    ws.cell(row=end + 2, column=2, value="Resumen").font = Font(bold=True)
    ws.cell(
        row=end + 2, column=3, value=f'=COUNTIF(E4:E{end},"✔")&" de "&ROWS(E4:E{end})&" correctas"'
    )
    ws.conditional_formatting.add(f"E4:E{end}", CellIsRule(operator="equal", formula=['"✔"'], fill=OK_FILL))
    ws.conditional_formatting.add(f"E4:E{end}", CellIsRule(operator="equal", formula=['"✘"'], fill=KO_FILL))
    for letter, width in zip("ABCDE", (6, 62, 22, 22, 12), strict=True):
        ws.column_dimensions[letter].width = width


# ---------------------------------------------------------------------------
# Ejercicios con tareas (matrices, LET/LAMBDA, agrupar)
# ---------------------------------------------------------------------------


def rng_p(col: str, n: int) -> str:
    return f"Pedidos!${col}$2:${col}${n + 1}"


def rng_c(col: str, m: int) -> str:
    return f"Clientes!${col}$2:${col}${m + 1}"


def exercise_matrices(p: pd.DataFrame, c: pd.DataFrame) -> Exercise:
    n, m = len(p), len(c)
    seg = c.set_index("customer_id")["segment"]
    top_order = p.loc[p["importe"].idxmax()]
    by_cust = p.groupby("customer_id")["importe"].sum()
    ex = Exercise(
        "matrices-dinamicas-y-buscarx",
        "Matrices dinámicas y BUSCARX",
        [
            "Objetivo: resolver con FILTRAR, UNICOS, ORDENARPOR y BUSCARX preguntas sobre 2.000 pedidos de Lumen.",
            "Hojas de datos: «Pedidos» (tabla tblPedidos) y «Clientes» (tabla tblClientes, con fichas incompletas a propósito).",
            "Pasos: 1) lee cada pregunta en «Trabajo»; 2) escribe tu fórmula en la celda amarilla; 3) mira «Comprobacion».",
            "Pista general: importe y coste están en euros; discount_pct es una fracción (0,2 = 20 %).",
        ],
    )
    ex.tasks = [
        Task(
            "¿Cuántos pedidos tienen un descuento superior al 20 %?",
            int((p["discount_pct"] > 0.2).sum()),
            f'COUNTIF({rng_p("E", n)},">0.2")',
            "=FILAS(FILTRAR(tblPedidos; tblPedidos[discount_pct]>0,2))",
            "FILTRAR / FILTER, FILAS / ROWS",
        ),
        Task(
            "¿Cuántos clientes distintos han hecho algún pedido?",
            int(p["customer_id"].nunique()),
            f"SUMPRODUCT(1/COUNTIF({rng_p('B', n)},{rng_p('B', n)}))",
            "=CONTARA(UNICOS(tblPedidos[customer_id]))",
            "UNICOS / UNIQUE, CONTARA / COUNTA",
        ),
        Task(
            "¿Qué segmento tiene el cliente del pedido de mayor importe? (si no hay ficha: «Sin segmento»)",
            str(seg.get(top_order["customer_id"], "Sin segmento")),
            f'IFERROR(INDEX({rng_c("B", m)},MATCH(INDEX({rng_p("B", n)},MATCH(MAX({rng_p("G", n)}),'
            f'{rng_p("G", n)},0)),{rng_c("A", m)},0)),"Sin segmento")',
            '=BUSCARX(BUSCARX(MAX(tblPedidos[importe]); tblPedidos[importe]; tblPedidos[customer_id]); '
            'tblClientes[customer_id]; tblClientes[segment]; "Sin segmento")',
            "BUSCARX / XLOOKUP, MAX",
        ),
        Task(
            "¿Cuántos pedidos pertenecen a clientes que NO tienen ficha en «Clientes»?",
            int((~p["customer_id"].isin(seg.index)).sum()),
            f"ROWS({rng_p('B', n)})-SUMPRODUCT(COUNTIF({rng_c('A', m)},{rng_p('B', n)}))",
            '=SUMA(--ESNOD(BUSCARX(tblPedidos[customer_id]; tblClientes[customer_id]; tblClientes[customer_id])))',
            "BUSCARX / XLOOKUP, ESNOD / ISNA",
        ),
        Task(
            "¿Cuál es el importe total del cliente que más ha comprado (suma de sus pedidos)?",
            round(float(by_cust.max()), 2),
            f"MAX(SUMIF({rng_p('B', n)},{rng_p('B', n)},{rng_p('G', n)}))",
            "=MAX(SUMAR.SI(tblPedidos[customer_id]; UNICOS(tblPedidos[customer_id]); tblPedidos[importe]))",
            "UNICOS / UNIQUE, SUMAR.SI / SUMIF, MAX",
        ),
    ]
    return ex


def exercise_let_lambda(p: pd.DataFrame, c: pd.DataFrame) -> Exercise:
    n = len(p)
    tot_imp, tot_cost = p["importe"].sum(), p["coste"].sum()
    hi = p[p["discount_pct"] >= 0.2]
    d = p["discount_pct"]
    scen = (p["importe"] * (1 - (d - 0.05).clip(lower=0)) / (1 - d)).sum()
    by_cust = p.groupby("customer_id")["importe"].sum()
    ex = Exercise(
        "let-lambda-escenarios",
        "LET, LAMBDA y escenarios",
        [
            "Objetivo: reescribir cálculos con LET, definir la función MARGEN con LAMBDA y construir un escenario.",
            "Hoja de datos: «Pedidos» (tabla tblPedidos, 2.000 pedidos). El importe ya es neto del descuento.",
            "Pasos: 1) en «Trabajo» escribe cada fórmula con LET; 2) crea la LAMBDA MARGEN en el Administrador de nombres;",
            "3) comprueba el resultado en «Comprobacion».",
        ],
    )
    ex.tasks = [
        Task(
            "Margen bruto total en % de los ingresos: (ingresos − coste) / ingresos (como fracción, p. ej. 0,35).",
            round(float((tot_imp - tot_cost) / tot_imp), 4),
            f"(SUM({rng_p('G', n)})-SUM({rng_p('H', n)}))/SUM({rng_p('G', n)})",
            "=LET(ing; SUMA(tblPedidos[importe]); cos; SUMA(tblPedidos[coste]); (ing-cos)/ing)",
            "LET",
            0.0005,
        ),
        Task(
            "Margen bruto en % (fracción) de los pedidos con descuento igual o superior al 20 %.",
            round(float((hi["importe"].sum() - hi["coste"].sum()) / hi["importe"].sum()), 4),
            f'(SUMIFS({rng_p("G", n)},{rng_p("E", n)},">=0.2")-SUMIFS({rng_p("H", n)},{rng_p("E", n)},">=0.2"))'
            f'/SUMIFS({rng_p("G", n)},{rng_p("E", n)},">=0.2")',
            '=LET(f; tblPedidos[discount_pct]>=0,2; ing; SUMA(FILTRAR(tblPedidos[importe]; f)); '
            "cos; SUMA(FILTRAR(tblPedidos[coste]; f)); (ing-cos)/ing)",
            "LET, FILTRAR / FILTER",
            0.0005,
        ),
        Task(
            "Crea la LAMBDA MARGEN(importe; coste) y escribe aquí el resultado de MARGEN(200; 130).",
            0.35,
            "(200-130)/200",
            "=MARGEN(200; 130)   con   MARGEN = LAMBDA(importe; coste; (importe-coste)/importe)",
            "LAMBDA",
            0.0005,
        ),
        Task(
            "Escenario: si cada descuento bajara 5 puntos (mínimo 0), ¿cuál sería el importe total? "
            "Nuevo importe = importe × (1 − nuevo descuento) / (1 − descuento actual).",
            round(float(scen), 2),
            f"SUMPRODUCT({rng_p('G', n)}*(1-({rng_p('E', n)}-0.05)*({rng_p('E', n)}>0.05))/(1-{rng_p('E', n)}))",
            "",
            "LET, SI / IF",
            0.5,
        ),
        Task(
            "¿Cuántos clientes han acumulado más de 500 € de importe?",
            int((by_cust > 500).sum()),
            f"SUMPRODUCT((SUMIF({rng_p('B', n)},{rng_p('B', n)},{rng_p('G', n)})>500)"
            f"/COUNTIF({rng_p('B', n)},{rng_p('B', n)}))",
            "=LET(cl; UNICOS(tblPedidos[customer_id]); tot; SUMAR.SI(tblPedidos[customer_id]; cl; "
            "tblPedidos[importe]); SUMA(--(tot>500)))",
            "LET, UNICOS / UNIQUE, SUMAR.SI / SUMIF",
        ),
    ]
    ex.tasks[3].modern = (
        "=LET(d; tblPedidos[discount_pct]; nd; SI(d>0,05; d-0,05; 0); "
        "SUMA(tblPedidos[importe]*(1-nd)/(1-d)))"
    )
    return ex


def exercise_agrupar(p: pd.DataFrame, c: pd.DataFrame) -> Exercise:
    n = len(p)
    by_ch = p.groupby("channel")["importe"].sum()
    ex = Exercise(
        "agrupar-y-pivotar",
        "Agrupar y pivotar",
        [
            "Objetivo: resumir 2.000 pedidos con AGRUPARPOR, PIVOTARPOR y una tabla dinámica y comprobar que coinciden.",
            "Hoja de datos: «Pedidos» (tabla tblPedidos). Estados: completado, devuelto_parcial y cancelado.",
            "Pasos: 1) crea en una hoja nueva una tabla dinámica de importe por canal y estado;",
            "2) escribe una fórmula PIVOTARPOR equivalente; 3) responde en «Trabajo» y revisa «Comprobacion».",
        ],
    )
    ex.tasks = [
        Task(
            "Importe total del canal «online».",
            round(float(by_ch["online"]), 2),
            f'SUMIFS({rng_p("G", n)},{rng_p("D", n)},"online")',
            '=SUMAR.SI.CONJUNTO(tblPedidos[importe]; tblPedidos[channel]; "online")   (o AGRUPARPOR)',
            "AGRUPARPOR / GROUPBY, SUMAR.SI.CONJUNTO / SUMIFS",
        ),
        Task(
            "Importe total del canal «store».",
            round(float(by_ch["store"]), 2),
            f'SUMIFS({rng_p("G", n)},{rng_p("D", n)},"store")',
            "=AGRUPARPOR(tblPedidos[channel]; tblPedidos[importe]; SUMA)",
            "AGRUPARPOR / GROUPBY",
        ),
        Task(
            "Número de pedidos cancelados del canal «online».",
            int(((p["channel"] == "online") & (p["status"] == "cancelado")).sum()),
            f'COUNTIFS({rng_p("D", n)},"online",{rng_p("F", n)},"cancelado")',
            "=PIVOTARPOR(tblPedidos[channel]; tblPedidos[status]; tblPedidos[order_id]; CONTARA)",
            "PIVOTARPOR / PIVOTBY, CONTAR.SI.CONJUNTO / COUNTIFS",
        ),
        Task(
            "Número total de pedidos con estado «devuelto_parcial».",
            int((p["status"] == "devuelto_parcial").sum()),
            f'COUNTIF({rng_p("F", n)},"devuelto_parcial")',
            '=CONTAR.SI(tblPedidos[status]; "devuelto_parcial")',
            "CONTAR.SI / COUNTIF",
        ),
        Task(
            "Comprobación de cuadre: suma de todos los totales por canal menos la suma del importe de la tabla "
            "(debe ser 0).",
            0,
            f'SUMIFS({rng_p("G", n)},{rng_p("D", n)},"online")+SUMIFS({rng_p("G", n)},{rng_p("D", n)},"store")'
            f"-SUM({rng_p('G', n)})",
            "=SUMA(tabla dinámica o AGRUPARPOR por canal) - SUMA(tblPedidos[importe])",
            "SUMA / SUM",
        ),
    ]
    return ex


def build_tasks_workbook(ex: Exercise, p: pd.DataFrame, c: pd.DataFrame, solution: bool) -> Workbook:
    wb = Workbook()
    write_instructions(wb, ex, solution)
    wsp = wb.create_sheet("Pedidos")
    write_table(wsp, p, "tblPedidos", date_cols=("order_date",))
    wsc = wb.create_sheet("Clientes")
    write_table(wsc, c, "tblClientes")
    ws = wb.create_sheet("Trabajo", 1)
    ws["A1"], ws["A1"].font = "Trabajo", Font(bold=True, size=14)
    heads = ["Nº", "Pregunta", "Tu respuesta", "Funciones sugeridas (ES / EN)"]
    if solution:
        heads.append("Fórmula moderna")
    for j, h in enumerate(heads, start=1):
        cell = ws.cell(row=3, column=j, value=h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    for i, t in enumerate(ex.tasks):
        r = 4 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=t.question).alignment = WRAP
        ans = ws.cell(row=r, column=3)
        ans.fill = INPUT_FILL
        if solution:
            ans.value = f"={t.classic}"
            modern = ws.cell(row=r, column=5, value=t.modern)
            modern.data_type = "s"  # texto, no fórmula
            modern.alignment = WRAP
        ws.cell(row=r, column=4, value=t.functions).alignment = WRAP
    for letter, width in zip("ABCDE", (6, 70, 24, 38, 80), strict=True):
        ws.column_dimensions[letter].width = width
    refs = [f"Trabajo!C{4 + i}" for i in range(len(ex.tasks))]
    write_check_sheet(
        wb,
        [t.question for t in ex.tasks],
        refs,
        [t.expected for t in ex.tasks],
        [t.tol for t in ex.tasks],
    )
    return wb


# ---------------------------------------------------------------------------
# Ejercicios de modelo con errores sembrados (auditoría)
# ---------------------------------------------------------------------------

MONTHS = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
COLS = "BCDEFGHIJKLM"  # meses
ROWS_MODEL = {
    "unidades": 5,
    "precio": 6,
    "brutos": 7,
    "dto": 8,
    "netos": 9,
    "coste": 10,
    "margen": 11,
    "fijos": 12,
    "ebit": 13,
    "impuestos": 14,
    "neto": 15,
    "margen_pct": 16,
}


def monthly_inputs() -> pd.DataFrame:
    orders = pd.read_parquet(DATA / "orders.parquet", columns=["order_id", "order_date", "status"])
    items = pd.read_parquet(DATA / "order_items.parquet")
    o = orders[(orders["order_date"].dt.year == 2025) & (orders["status"] == "completado")]
    it = items.merge(o[["order_id", "order_date"]], on="order_id")
    it["mes"] = it["order_date"].dt.month
    g = it.groupby("mes").agg(unidades=("quantity", "sum"), brutos=("unit_price", "mean"))
    g["precio"] = g["brutos"].round(2)
    return g[["unidades", "precio"]]


def model_values(inputs: pd.DataFrame, dto: float, coste_u: float, fijos: float, tax: float) -> dict:
    u, pr = inputs["unidades"].to_numpy(), inputs["precio"].to_numpy()
    brutos = u * pr
    dtos = brutos * dto
    netos = brutos - dtos
    coste = u * coste_u
    margen = netos - coste
    ebit = margen - fijos
    imp = np.maximum(ebit, 0) * tax
    neto = ebit - imp
    return {
        "u": u,
        "brutos": brutos,
        "dto": dtos,
        "netos": netos,
        "coste": coste,
        "margen": margen,
        "ebit": ebit,
        "imp": imp,
        "neto": neto,
        "mpct": margen / netos,
    }


def write_model(wb: Workbook, inputs: pd.DataFrame, assumptions: dict, errors: bool) -> None:
    ws = wb.create_sheet("Supuestos")
    ws["A1"], ws["A1"].font = "Supuestos", Font(bold=True, size=14)
    rows = [
        ("Descuento medio sobre el precio bruto", assumptions["dto"], "0%"),
        ("Coste unitario (€ por unidad)", assumptions["coste_u"], "#,##0.00"),
        ("Gastos fijos mensuales (€)", assumptions["fijos"], "#,##0"),
        ("Tipo impositivo", assumptions["tax"], "0%"),
    ]
    for i, (lab, val, fmt) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=lab)
        cell = ws.cell(row=i, column=2, value=val)
        cell.number_format, cell.fill = fmt, INPUT_FILL
    ws.column_dimensions["A"].width = 44
    if errors:
        ws["B5"] = 21  # error 6: 21 en lugar de 21 %
        ws["B5"].number_format = "General"

    m = wb.create_sheet("Modelo")
    m["A1"], m["A1"].font = "Modelo de resultados 2025 (euros)", Font(bold=True, size=14)
    labels = {
        5: "Unidades",
        6: "Precio bruto medio (€)",
        7: "Ingresos brutos",
        8: "Descuento",
        9: "Ingresos netos",
        10: "Coste de ventas",
        11: "Margen bruto",
        12: "Gastos fijos",
        13: "EBIT",
        14: "Impuestos",
        15: "Resultado neto",
        16: "Margen bruto % s/ ingresos netos",
    }
    m.cell(row=4, column=1, value="Concepto").font = Font(bold=True)
    for j, mes in enumerate(MONTHS):
        m.cell(row=4, column=2 + j, value=mes).font = Font(bold=True)
    m.cell(row=4, column=14, value="Total 2025").font = Font(bold=True)
    for r, lab in labels.items():
        m.cell(row=r, column=1, value=lab)
    for j, col in enumerate(COLS):
        u, pr = inputs["unidades"].iloc[j], inputs["precio"].iloc[j]
        m[f"{col}5"], m[f"{col}6"] = float(u), float(pr)
        m[f"{col}5"].fill = m[f"{col}6"].fill = INPUT_FILL
        m[f"{col}7"] = f"={col}5*{col}6"
        m[f"{col}8"] = f"={col}7*Supuestos!$B$2"
        m[f"{col}9"] = f"={col}7-{col}8"
        m[f"{col}10"] = f"={col}5*Supuestos!$B$3"
        m[f"{col}11"] = f"={col}9-{col}10"
        m[f"{col}12"] = "=Supuestos!$B$4"
        m[f"{col}13"] = f"={col}11-{col}12"
        m[f"{col}14"] = f"=MAX(0,{col}13)*Supuestos!$B$5"
        m[f"{col}15"] = f"={col}13-{col}14"
        m[f"{col}16"] = f"={col}11/{col}9"
    for r in (5, 7, 8, 9, 10, 11, 12, 13, 14, 15):
        m[f"N{r}"] = f"=SUM(B{r}:M{r})"
    m["N6"] = "=N7/N5"
    m["N16"] = "=N11/N9"
    for r in range(5, 16):
        for col in [*COLS, "N"]:
            m[f"{col}{r}"].number_format = "#,##0.00" if r == 6 else "#,##0"
    for col in [*COLS, "N"]:
        m[f"{col}16"].number_format = "0.0%"
    m.column_dimensions["A"].width = 36
    if errors:
        m["N9"] = "=SUM(B9:L9)"  # error 1: rango corto
        m["H10"] = round(float(inputs["unidades"].iloc[6]) * assumptions["coste_u"] * 1.08)  # error 2: valor pegado (julio)
        m["F8"] = "=F7*Supuestos!F2"  # error 3: referencia no anclada
        m["D5"] = str(int(inputs["unidades"].iloc[2]))  # error 4: número como texto
        m["J16"] = "=J11/J10"  # error 5: margen sobre coste
        m["N13"] = "=SUM(B13:M13)+B13"  # error 7: doble recuento de enero


def exercise_model(kind: str) -> tuple[Exercise, dict]:
    inputs = monthly_inputs()
    avg_net = float((inputs["unidades"] * inputs["precio"]).mean()) * 0.92
    a = {
        "dto": 0.08,
        "coste_u": round(float(inputs["precio"].mean()) * 0.55, 2),
        "fijos": round(avg_net * 0.15, -3),
        "tax": 0.21,
    }
    v = model_values(inputs, a["dto"], a["coste_u"], a["fijos"], a["tax"])
    ex = Exercise(
        "auditoria-de-un-modelo",
        "Auditoría de un modelo con errores",
        [
            "Objetivo: encontrar y corregir SIETE errores sembrados en el modelo de resultados de 2025.",
            "Cómo trabajar: usa Fórmulas > Auditoría de fórmulas (precedentes, dependientes, evaluar) y Ctrl+` para ver fórmulas.",
            "Corrige las fórmulas (no pegues valores). Cuando todos los controles de «Comprobacion» estén en ✔, has terminado.",
            "Después compara con el libro de solución, donde la hoja «Errores» describe cada fallo.",
        ],
        kind="model",
    )
    checks = [
        ("Ingresos netos totales (Modelo!N9)", "Modelo!N9", round(float(v["netos"].sum()), 2), 0.5),
        ("Coste de ventas total (Modelo!N10)", "Modelo!N10", round(float(v["coste"].sum()), 2), 0.5),
        ("Descuento de mayo (Modelo!F8)", "Modelo!F8", round(float(v["dto"][4]), 2), 0.5),
        ("Unidades totales (Modelo!N5)", "Modelo!N5", float(v["u"].sum()), 0.5),
        ("Margen bruto % de septiembre (Modelo!J16)", "Modelo!J16", round(float(v["mpct"][8]), 4), 0.0005),
        ("Impuestos totales (Modelo!N14)", "Modelo!N14", round(float(v["imp"].sum()), 2), 0.5),
        ("EBIT total (Modelo!N13)", "Modelo!N13", round(float(v["ebit"].sum()), 2), 0.5),
        ("Coste de ventas de julio (Modelo!H10)", "Modelo!H10", round(float(v["coste"][6]), 2), 0.5),
    ]
    return ex, {"inputs": inputs, "assumptions": a, "checks": checks, "values": v}


MODEL_ERRORS = [
    ("Modelo!N9", "El total de ingresos netos suma solo hasta noviembre (B9:L9): falta diciembre.", "=SUM(B9:M9)"),
    ("Modelo!H10", "El coste de julio es un número pegado (y desactualizado) en lugar de la fórmula.", "=H5*Supuestos!$B$3"),
    ("Modelo!F8", "El descuento de mayo apunta a Supuestos!F2 (vacía) porque la referencia no está anclada.", "=F7*Supuestos!$B$2"),
    ("Modelo!D5", "Las unidades de marzo están guardadas como texto: SUM las ignora en el total.", "Convertir a número"),
    ("Modelo!J16", "El margen de septiembre se divide entre el coste, no entre los ingresos netos.", "=J11/J9"),
    ("Supuestos!B5", "El tipo impositivo está escrito como 21 en lugar de 21 % (0,21).", "0,21"),
    ("Modelo!N13", "El EBIT total suma además el de enero dos veces.", "=SUM(B13:M13)"),
]


def build_model_workbook(solution: bool) -> Workbook:
    ex, ctx = exercise_model("auditoria")
    wb = Workbook()
    write_instructions(wb, ex, solution)
    write_model(wb, ctx["inputs"], ctx["assumptions"], errors=not solution)
    ch = ctx["checks"]
    write_check_sheet(wb, [c[0] for c in ch], [c[1] for c in ch], [c[2] for c in ch], [c[3] for c in ch])
    if solution:
        ws = wb.create_sheet("Errores")
        for j, h in enumerate(["Celda", "Qué estaba mal", "Corrección"], start=1):
            cell = ws.cell(row=1, column=j, value=h)
            cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        for i, (cell_ref, what, fix) in enumerate(MODEL_ERRORS, start=2):
            ws.cell(row=i, column=1, value=cell_ref)
            ws.cell(row=i, column=2, value=what).alignment = WRAP
            c = ws.cell(row=i, column=3, value=fix)
            c.data_type = "s"
        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 90
        ws.column_dimensions["C"].width = 28
    return wb


# --- Verificar el trabajo de una IA ------------------------------------------------------


def ai_ctx() -> dict:
    p, _ = load_sample()
    c = pd.read_parquet(DATA / "customers.parquet")[["customer_id", "segment"]]
    o = pd.read_parquet(DATA / "orders.parquet", columns=["order_id", "customer_id", "order_date", "status", "discount_pct"])
    items = pd.read_parquet(DATA / "order_items.parquet")
    prods = pd.read_parquet(DATA / "products.parquet")[["product_id", "cost"]]
    o = o[(o["order_date"].dt.year == 2025) & (o["status"] == "completado")].merge(c, on="customer_id")
    it = items.merge(o, on="order_id").merge(prods, on="product_id")
    it["importe"], it["coste"] = it["quantity"] * it["unit_price"], it["quantity"] * it["cost"]
    g = it.groupby("segment").agg(importe=("importe", "sum"), coste=("coste", "sum"), dto=("discount_pct", "mean"))
    g = g.loc[["nuevo", "ocasional", "habitual", "vip"]]
    g["importe"] = (g["importe"] / 1000).round(0)  # en miles de euros
    g["coste"] = (g["coste"] / 1000).round(0)
    g["dto"] = g["dto"].round(3)
    return {"g": g, "obj": 0.12}


def ai_values(g: pd.DataFrame, obj: float) -> dict:
    esc = g["importe"] * (1 - obj) / (1 - g["dto"])
    marg = esc - g["coste"]
    return {"esc": esc, "marg": marg, "mpct": marg / esc}


def build_ai_workbook(solution: bool) -> Workbook:
    ctx = ai_ctx()
    g, obj = ctx["g"], ctx["obj"]
    v = ai_values(g, obj)
    wb = Workbook()
    ex = Exercise(
        "verificar-el-trabajo-de-una-ia",
        "Verificar el trabajo de una IA",
        [
            "Situación: un asistente de IA actualizó la hipótesis de descuento de un modelo de previsión por segmento.",
            "El resultado «tiene buena pinta», pero el asistente cometió CUATRO errores plausibles.",
            "Objetivo: aplicar el protocolo de cuatro controles (cuadre, muestreo, fórmulas visibles, casos límite),",
            "encontrar los errores, corregirlos y dejar todos los controles de «Comprobacion» en ✔.",
            "Importes en miles de euros (2025, pedidos completados).",
        ],
        kind="model",
    )
    write_instructions(wb, ex, solution)
    s = wb.create_sheet("Supuestos")
    s["A1"], s["A1"].font = "Supuestos", Font(bold=True, size=14)
    s["A2"], s["B2"] = "Descuento objetivo", obj
    s["B2"].number_format, s["B2"].fill = "0%", INPUT_FILL
    s.column_dimensions["A"].width = 30
    if not solution:
        s["B2"] = 12  # error 1: 12 en lugar de 12 %
        s["B2"].number_format = "General"

    w = wb.create_sheet("Prevision")
    heads = ["Segmento", "Ingresos 2025", "Coste 2025", "Descuento actual", "Ingresos escenario", "Coste escenario", "Margen", "Margen %"]
    for j, h in enumerate(heads, start=1):
        cell = w.cell(row=1, column=j, value=h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    for i, (seg, row) in enumerate(g.iterrows(), start=2):
        w.cell(row=i, column=1, value=seg)
        for j, key in enumerate(("importe", "coste", "dto"), start=2):
            cell = w.cell(row=i, column=j, value=float(row[key]))
            cell.fill = INPUT_FILL
        w[f"E{i}"] = f"=B{i}*(1-Supuestos!$B$2)/(1-D{i})"
        w[f"F{i}"] = f"=C{i}"
        w[f"G{i}"] = f"=E{i}-F{i}"
        w[f"H{i}"] = f"=G{i}/E{i}"
        w[f"H{i}"].number_format = "0.0%"
        w[f"D{i}"].number_format = "0.0%"
    w["A6"] = "Total"
    for col in "BCEFG":
        w[f"{col}6"] = f"=SUM({col}2:{col}5)"
    w["H6"] = "=G6/E6"
    w["H6"].number_format = "0.0%"
    for col in "ABCDEFGH":
        w[f"{col}6"].font = Font(bold=True)
    if not solution:
        w["E6"] = "=SUM(E2:E4)"  # error 2: rango truncado
        w["H3"] = "=G3/B3"  # error 3: margen % sobre ingresos base
        w["G4"] = round(float(v["marg"].iloc[2]) * 0.93, 1)  # error 4: valor pegado desactualizado
    w.column_dimensions["A"].width = 14
    for col in "BCDEFGH":
        w.column_dimensions[col].width = 18
    checks = [
        ("Ingresos escenario total (Prevision!E6)", "Prevision!E6", round(float(v["esc"].sum()), 2), 0.5),
        ("Margen % del segmento ocasional (Prevision!H3)", "Prevision!H3", round(float(v["mpct"].iloc[1]), 4), 0.0005),
        ("Margen % del segmento habitual (Prevision!H4)", "Prevision!H4", round(float(v["mpct"].iloc[2]), 4), 0.0005),
        ("Margen del segmento habitual (Prevision!G4)", "Prevision!G4", round(float(v["marg"].iloc[2]), 2), 0.5),
        ("Margen total (Prevision!G6)", "Prevision!G6", round(float(v["marg"].sum()), 2), 0.5),
        ("Margen % total (Prevision!H6)", "Prevision!H6", round(float(v["marg"].sum() / v["esc"].sum()), 4), 0.0005),
    ]
    write_check_sheet(wb, [c[0] for c in checks], [c[1] for c in checks], [c[2] for c in checks], [c[3] for c in checks])
    if solution:
        e = wb.create_sheet("Errores")
        for j, h in enumerate(["Celda", "Qué estaba mal", "Cómo se detecta"], start=1):
            cell = e.cell(row=1, column=j, value=h)
            cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        rows = [
            ("Supuestos!B2", "El descuento está escrito como 12 en lugar de 12 % (0,12).", "Los ingresos del escenario son negativos o absurdos: caso límite."),
            ("Prevision!E6", "El total suma E2:E4 y deja fuera el segmento vip.", "Cuadre: el total no coincide con la suma de las filas."),
            ("Prevision!H3", "El margen % del segmento ocasional se divide entre los ingresos base (B3).", "Fórmulas visibles: la fila no sigue el patrón de las demás."),
            ("Prevision!G4", "El margen del segmento habitual es un valor pegado, no una fórmula.", "Ir a Especial > Constantes dentro del bloque de fórmulas."),
        ]
        for i, row in enumerate(rows, start=2):
            for j, val in enumerate(row, start=1):
                e.cell(row=i, column=j, value=val).alignment = WRAP
        e.column_dimensions["A"].width = 16
        e.column_dimensions["B"].width = 70
        e.column_dimensions["C"].width = 60
    return wb


# ---------------------------------------------------------------------------


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    p, c = load_sample()
    builders = {
        "matrices-dinamicas-y-buscarx": lambda: exercise_matrices(p, c),
        "let-lambda-escenarios": lambda: exercise_let_lambda(p, c),
        "agrupar-y-pivotar": lambda: exercise_agrupar(p, c),
    }
    for wid, build in builders.items():
        ex = build()
        for sol in (False, True):
            wb = build_tasks_workbook(ex, p, c, sol)
            wb.save(OUT / f"{wid}{'-solucion' if sol else ''}.xlsx")
    for sol in (False, True):
        build_model_workbook(sol).save(OUT / f"auditoria-de-un-modelo{'-solucion' if sol else ''}.xlsx")
        build_ai_workbook(sol).save(OUT / f"verificar-el-trabajo-de-una-ia{'-solucion' if sol else ''}.xlsx")
    print(f"Libros generados en {OUT}: {len(list(OUT.glob('*.xlsx')))}")


if __name__ == "__main__":
    sys.exit(main())
